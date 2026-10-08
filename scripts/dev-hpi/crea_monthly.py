"""Adapter verified against CREA's September 2026 non-adjusted monthly workbook."""

from pathlib import Path

HOUSING_COLUMNS = {
    "Composite": "Composite",
    "Single Family": "Single_Family",
    "Townhouse": "Townhouse",
    "Apartment": "Apartment",
}


def monthly_records(path):
    """Read market sheets without merging seasonally adjusted or non-monthly data."""
    from openpyxl import load_workbook

    if Path(path).name != "Not Seasonally Adjusted (M).xlsx":
        raise ValueError("Use the official Not Seasonally Adjusted (M).xlsx workbook")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in workbook:
            rows = sheet.iter_rows(values_only=True)
            headers = list(next(rows))
            allowed = {"Date"}
            for prefix in (*HOUSING_COLUMNS.values(), "One_Storey", "Two_Storey"):
                allowed.update((prefix + "_HPI", prefix + "_Benchmark"))
            if not headers or headers[0] != "Date" or len(set(headers)) != len(headers):
                raise ValueError("Official monthly headers changed")
            if not set(headers).issubset(allowed):
                raise ValueError("Unexpected columns or adjusted series")
            if not {"Composite_HPI", "Composite_Benchmark"}.issubset(headers):
                raise ValueError("Missing composite series")
            for cells in rows:
                if all(value is None for value in cells):
                    continue
                for housing, prefix in HOUSING_COLUMNS.items():
                    record = {
                        "market": sheet.title.replace("_", " "),
                        "housing_type": housing,
                        "month": cells[0],
                    }
                    for metric, suffix in (
                        ("benchmark_price", "_Benchmark"),
                        ("hpi_index", "_HPI"),
                    ):
                        header = prefix + suffix
                        record[metric] = (
                            cells[headers.index(header)] if header in headers else None
                        )
                    yield record
    finally:
        workbook.close()

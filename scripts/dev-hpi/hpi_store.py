"""Private, local HPI storage. Never writes a workbook or sends data to a service."""

import argparse
from contextlib import closing
import datetime as dt
import hashlib
import json
import math
from pathlib import Path
import re
import sqlite3

SOURCE = "https://www.crea.ca/housing-market-stats/mls-home-price-index/hpi-tool/"
TYPES = ("Composite", "Single Family", "Townhouse", "Apartment")
METRICS = (
    "benchmark_price",
    "hpi_index",
    "change_1m",
    "change_6m",
    "change_1y",
    "change_3y",
    "change_5y",
)


def private_path(value):
    """Reject paths inside any Git checkout, including linked worktrees."""
    path = Path(value).expanduser().resolve()
    if any((parent / ".git").exists() for parent in (path, *path.parents)):
        raise ValueError("HPI source and database must be outside Git checkouts")
    return path


def month_key(value):
    """Normalize an explicit month without inventing its year."""
    if isinstance(value, (dt.datetime, dt.date)):
        return value.strftime("%Y-%m")
    text = str(value).strip()
    if not re.fullmatch(r"\d{4}-\d{2}(?:-\d{2})?", text):
        raise ValueError("Month must be an Excel date or YYYY-MM")
    return dt.date.fromisoformat(text if len(text) == 10 else text + "-01").strftime(
        "%Y-%m"
    )


def numeric(value, percentage=False):
    """Keep missing values null and preserve percentage points (not fractions)."""
    if value is None or str(value).strip().lower() in (
        "",
        "n/a",
        "na",
        "--",
        "-",
        "..",
        "...",
    ):
        return None
    text = str(value).strip().replace(",", "").replace("$", "")
    if text.endswith("%"):
        if not percentage:
            raise ValueError("Unexpected percentage in level column")
        text = text[:-1]
    number = float(text)
    if not math.isfinite(number):
        raise ValueError("Non-finite source number")
    return number


def normalize(record):
    """Validate source fields; missing levels and change percentages remain null."""
    market = str(record.get("market") or "").strip()
    housing = str(record.get("housing_type") or "").strip()
    if not market or len(market) > 200 or housing not in TYPES:
        raise ValueError("Missing market or unsupported housing type")
    result = {
        "market": market,
        "housing_type": housing,
        "month": month_key(record.get("month")),
    }
    for metric in METRICS:
        result[metric] = numeric(record.get(metric), metric.startswith("change_"))
        if (
            metric in ("benchmark_price", "hpi_index")
            and result[metric] is not None
            and result[metric] <= 0
        ):
            raise ValueError(
                "Benchmark and index levels must be positive when supplied"
            )
    return result


def workbook_records(path, layout):
    """Read explicit long-format sheet mappings; fail on schema drift, never guess.

    Mapping: sheets=[{name, header_row, columns: {canonical_field: source_header},
    percentage_units: 'points'|'fraction'}]. openpyxl is used only for reading.
    An official-source adapter must be verified against the downloaded workbook.
    """
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in layout["sheets"]:
            rows = workbook[sheet["name"]].iter_rows(values_only=True)
            for _ in range(sheet.get("header_row", 1) - 1):
                next(rows)
            headers = [
                str(item).strip() if item is not None else "" for item in next(rows)
            ]
            columns = sheet["columns"]
            if not {"market", "housing_type", "month"}.issubset(columns) or not set(
                columns
            ).issubset({"market", "housing_type", "month", *METRICS}):
                raise ValueError("Invalid column mapping")
            if any(headers.count(label) != 1 for label in columns.values()):
                raise ValueError("Mapped header missing or duplicated")
            units = sheet.get("percentage_units", "points")
            if units not in ("points", "fraction"):
                raise ValueError("Unknown percentage units")
            indexes = {key: headers.index(label) for key, label in columns.items()}
            for cells in rows:
                if all(cell is None for cell in cells):
                    continue
                record = {key: cells[index] for key, index in indexes.items()}
                if units == "fraction":
                    for key in METRICS[2:]:
                        value = record.get(key)
                        if value is not None:
                            if isinstance(value, str) and value.strip().endswith("%"):
                                raise ValueError(
                                    "Fraction mapping cannot contain percent text"
                                )
                            parsed = numeric(value, True)
                            record[key] = None if parsed is None else parsed * 100
                yield normalize(record)
    finally:
        workbook.close()


def import_records(database, records, digest, synthetic=False, derive_changes=False):
    """Atomically upsert monthly observations; reject ambiguous duplicate keys."""
    database = private_path(database)
    database.parent.mkdir(parents=True, exist_ok=True)
    seen = set()
    prepared = []
    for record in records:
        row = normalize(record)
        key = tuple(row[field] for field in ("market", "housing_type", "month"))
        if key in seen:
            raise ValueError("Duplicate market/type/month in import")
        seen.add(key)
        prepared.append(row)
    if not prepared:
        raise ValueError("No observations found")
    with closing(sqlite3.connect(database)) as db, db:
        db.execute(
            "CREATE TABLE IF NOT EXISTS hpi_settings (change_basis TEXT NOT NULL)"
        )
        basis = (
            "derived_hpi_exact_months" if derive_changes else "source_percentage_points"
        )
        previous_basis = db.execute("SELECT change_basis FROM hpi_settings").fetchone()
        if previous_basis and previous_basis[0] != basis:
            raise ValueError("Use separate databases for different series/change bases")
        if not previous_basis:
            db.execute("INSERT INTO hpi_settings VALUES (?)", (basis,))
        db.execute("""CREATE TABLE IF NOT EXISTS hpi_benchmarks (
            market TEXT NOT NULL, housing_type TEXT NOT NULL, month TEXT NOT NULL,
            benchmark_price REAL, hpi_index REAL, change_1m REAL, change_6m REAL,
            change_1y REAL, change_3y REAL, change_5y REAL,
            source_sha256 TEXT NOT NULL, imported_at TEXT NOT NULL,
            synthetic INTEGER NOT NULL,
            PRIMARY KEY (market, housing_type, month))""")
        # A database is either official/private or synthetic: never silently mix it.
        kinds = db.execute("SELECT DISTINCT synthetic FROM hpi_benchmarks").fetchall()
        if kinds and kinds != [(int(synthetic),)]:
            raise ValueError("Use separate databases for synthetic and official data")
        fields = ("market", "housing_type", "month", *METRICS)
        stamp = dt.datetime.now(dt.timezone.utc).isoformat()
        updates = ", ".join(
            f"{field}=excluded.{field}"
            for field in (*METRICS, "source_sha256", "imported_at", "synthetic")
        )
        db.executemany(
            f"INSERT INTO hpi_benchmarks VALUES ({','.join('?' for _ in range(13))}) "
            f"ON CONFLICT(market,housing_type,month) DO UPDATE SET {updates}",
            [
                tuple(row[field] for field in fields) + (digest, stamp, int(synthetic))
                for row in prepared
            ],
        )
        if derive_changes:
            derive_monthly_changes(db)
    return {
        "imported_rows": len(prepared),
        "markets": len({row["market"] for row in prepared}),
    }


def derive_monthly_changes(db):
    """Recompute from exact HPI month lags, including existing imported history.

    Percent change = (current index / prior index - 1) * 100. Missing current
    or exact prior month stays null; never substitute the nearest available month.
    Recomputing history also propagates revisions from newly appended workbooks.
    """
    series = {
        (market, housing, month): index
        for market, housing, month, index in db.execute(
            "SELECT market, housing_type, month, hpi_index FROM hpi_benchmarks"
        )
    }
    updates = []
    for (market, housing, month), current in series.items():
        year, month_number = map(int, month.split("-"))
        serial = year * 12 + month_number - 1
        changes = []
        for lag in (1, 6, 12, 36, 60):
            prior_year, prior_month = divmod(serial - lag, 12)
            prior = series.get(
                (market, housing, f"{prior_year:04d}-{prior_month + 1:02d}")
            )
            changes.append((current / prior - 1) * 100 if current and prior else None)
        updates.append((*changes, market, housing, month))
    db.executemany(
        "UPDATE hpi_benchmarks SET change_1m=?, change_6m=?, change_1y=?, "
        "change_3y=?, change_5y=? WHERE market=? AND housing_type=? AND month=?",
        updates,
    )


def query(database, market=None, housing_type=None, month=None):
    """Read a private database without creating a missing file."""
    database = private_path(database)
    with closing(sqlite3.connect(database.as_uri() + "?mode=ro", uri=True)) as db:
        db.row_factory = sqlite3.Row
        choices = [
            dict(row)
            for row in db.execute(
                "SELECT market, housing_type, month FROM hpi_benchmarks "
                "ORDER BY market, housing_type, month DESC"
            )
        ]
        result = None
        if market and housing_type and month:
            row = db.execute(
                "SELECT * FROM hpi_benchmarks WHERE market=? AND housing_type=? AND month=?",
                (market, housing_type, month_key(month)),
            ).fetchone()
            result = dict(row) if row else None
            if result:
                basis = db.execute("SELECT change_basis FROM hpi_settings").fetchone()
                result["change_basis"] = (
                    basis[0] if basis else "source_percentage_points"
                )
                result["series_adjustment"] = (
                    "not_seasonally_adjusted"
                    if basis and basis[0] == "derived_hpi_exact_months"
                    else "unspecified"
                )
        return {"choices": choices, "benchmark": result, "source_url": SOURCE}


def main():
    """CLI prints row counts on imports, values only for the private read bridge."""
    parser = argparse.ArgumentParser()
    parser.add_argument("operation", choices=("import", "query"))
    parser.add_argument("--database", required=True)
    parser.add_argument("--workbook")
    parser.add_argument("--layout")
    parser.add_argument("--crea-monthly", action="store_true")
    parser.add_argument("--market")
    parser.add_argument("--housing-type")
    parser.add_argument("--month")
    args = parser.parse_args()
    if args.operation == "import":
        source = private_path(args.workbook)
        if args.crea_monthly:
            from crea_monthly import monthly_records

            records = monthly_records(source)
        else:
            layout = json.loads(Path(args.layout).read_text(encoding="utf-8"))
            records = workbook_records(source, layout)
        digest = hashlib.sha256(source.read_bytes()).hexdigest()
        result = import_records(
            args.database, records, digest, derive_changes=args.crea_monthly
        )
    else:
        result = query(args.database, args.market, args.housing_type, args.month)
    print(json.dumps(result))


if __name__ == "__main__":
    main()
